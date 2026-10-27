from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
OUTPUT = ROOT / "deliverables" / "比亚迪股票价格走势分析_2023-2026.xlsx"

INK = "17212B"
RED = "C73A42"
GREEN = "15805D"
PAPER = "F4F4EF"
GRID = "D8DAD4"
BLUE = "2E5C8A"
GOLD = "A8752A"


def style_header(row) -> None:
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")


def set_widths(worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width


def title_block(worksheet, title: str, subtitle: str, width: int = 12) -> None:
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    title_cell = worksheet.cell(1, 1, title)
    title_cell.font = Font(size=18, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor=INK)
    title_cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 34

    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    subtitle_cell = worksheet.cell(2, 1, subtitle)
    subtitle_cell.font = Font(size=10, color="52616F")
    subtitle_cell.fill = PatternFill("solid", fgColor=PAPER)
    subtitle_cell.alignment = Alignment(vertical="center", wrap_text=True)
    worksheet.row_dimensions[2].height = 28


def add_data_table(worksheet, start_row: int, end_row: int, end_col: int, name: str) -> None:
    ref = f"A{start_row}:{get_column_letter(end_col)}{end_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def add_positive_negative_rules(worksheet, cell_range: str) -> None:
    positive_fill = PatternFill("solid", fgColor="FBE4E4")
    negative_fill = PatternFill("solid", fgColor="DDF2EA")
    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(operator="greaterThan", formula=["0"], fill=positive_fill),
    )
    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(operator="lessThan", formula=["0"], fill=negative_fill),
    )


def build_notes(workbook: Workbook, summary: dict) -> None:
    worksheet = workbook.active
    worksheet.title = "任务说明"
    worksheet.sheet_view.showGridLines = False
    title_block(
        worksheet,
        "比亚迪（002594.SZ）股票价格走势分析与可视化",
        "教材 P92 实战演练 · 2023-01-03 至 2026-10-27 · 后复权日频行情",
        width=8,
    )

    metadata = summary["metadata"]
    metrics = summary["keyMetrics"]
    validation = summary["validation"]
    evidence = [
        ("数据区间", f"{metadata['startDate']} 至 {metadata['endDate']}"),
        ("交易日数", f"{metadata['tradingDays']} 个交易日"),
        ("价格口径", metadata["priceBasis"]),
        ("2023 年原始数据", metadata["sourceOriginal"]),
        ("2024-2026 扩展数据", metadata["sourceExtension"]),
        (
            "口径核验",
            f"2023 年全部 {validation['overlapRows']} 行逐字段差异为 0，"
            "交易日历完全一致",
        ),
        ("最新收盘", f"{metrics['latestClose']:.2f} 元"),
        ("区间收益", f"{metrics['periodReturnPct']:.2f}%"),
        ("最大回撤", f"{metrics['maxDrawdownPct']:.2f}%"),
    ]
    for row_index, (label, value) in enumerate(evidence, start=4):
        worksheet.cell(row_index, 1, label).font = Font(bold=True, color=INK)
        worksheet.merge_cells(
            start_row=row_index, start_column=2, end_row=row_index, end_column=8
        )
        worksheet.cell(row_index, 2, value).alignment = Alignment(wrap_text=True)
        worksheet.row_dimensions[row_index].height = 24

    worksheet.cell(14, 1, "指标口径").font = Font(size=12, bold=True, color=INK)
    methodology = summary["methodology"]
    method_rows = [
        ("收益率", methodology["return"]),
        ("年化波动率", methodology["volatility"]),
        ("最大回撤", methodology["drawdown"]),
        ("均线", methodology["ma"]),
        ("RSI", methodology["rsi"]),
        ("年度收益", methodology["annualReturn"]),
        (
            "单位",
            "价格：元；成交量：手；成交额：亿元；涨跌幅、振幅、换手率：%。",
        ),
        (
            "说明",
            "后复权考虑了分红送转与除权除息，适合观察连续价格趋势。"
            "分析为描述性统计，不构成投资建议。",
        ),
    ]
    for row_index, (label, value) in enumerate(method_rows, start=15):
        worksheet.cell(row_index, 1, label).font = Font(bold=True, color=INK)
        worksheet.merge_cells(
            start_row=row_index, start_column=2, end_row=row_index, end_column=8
        )
        worksheet.cell(row_index, 2, value).alignment = Alignment(wrap_text=True)
        worksheet.row_dimensions[row_index].height = 28

    worksheet.cell(24, 1, "工作表导航").font = Font(size=12, bold=True, color=INK)
    guide = [
        ("完整行情", "922 个交易日的 OHLC、成交量额、换手率与指标"),
        ("年度统计", "年度收益、高低价、成交额与日均换手率"),
        ("月度统计", "月末价格、月度收益、成交额与换手率"),
        ("指标摘要", "全区间关键指标与风险口径"),
        ("收盘走势", "收盘价与 MA20 / MA60 / MA120 原生折线图"),
        ("成交量", "日成交量原生柱状图"),
    ]
    for row_index, (sheet, description) in enumerate(guide, start=25):
        worksheet.cell(row_index, 1, sheet).font = Font(bold=True, color=BLUE)
        worksheet.merge_cells(
            start_row=row_index, start_column=2, end_row=row_index, end_column=8
        )
        worksheet.cell(row_index, 2, description)

    set_widths(worksheet, [20, 18, 18, 18, 18, 18, 18, 18])
    worksheet.freeze_panes = "A4"


def build_market_sheet(workbook: Workbook, frame: pd.DataFrame) -> None:
    worksheet = workbook.create_sheet("完整行情")
    worksheet.sheet_view.showGridLines = False
    columns = [
        "日期",
        "开盘",
        "收盘",
        "最高",
        "最低",
        "成交量",
        "成交额",
        "振幅",
        "涨跌幅",
        "涨跌额",
        "换手率",
        "MA20",
        "MA60",
        "MA120",
        "回撤",
        "RSI14",
        "MACD_DIF",
        "MACD_DEA",
        "MACD柱",
    ]
    worksheet.append(columns)
    style_header(worksheet[1])
    for row in frame[columns].itertuples(index=False, name=None):
        worksheet.append(list(row))

    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        row[0].number_format = "yyyy-mm-dd"
        for cell in row[1:]:
            cell.number_format = "0.0000"
        for index in (7, 8, 9, 10, 14):
            row[index].number_format = "0.00"
    add_data_table(worksheet, 1, worksheet.max_row, len(columns), "MarketData")
    add_positive_negative_rules(
        worksheet, f"I2:I{worksheet.max_row}"
    )
    worksheet.freeze_panes = "B2"
    worksheet.auto_filter.ref = f"A1:S{worksheet.max_row}"
    set_widths(worksheet, [12] + [11] * (len(columns) - 1))
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 2


def build_annual_sheet(workbook: Workbook, annual: pd.DataFrame) -> None:
    worksheet = workbook.create_sheet("年度统计")
    worksheet.sheet_view.showGridLines = False
    title_block(
        worksheet,
        "年度表现",
        "红涨绿跌；2026 年为截至 2026-10-27 的年内表现。",
        width=9,
    )
    headers = list(annual.columns)
    start_row = 4
    for column_index, header in enumerate(headers, start=1):
        worksheet.cell(start_row, column_index, header)
    style_header(worksheet[start_row])
    for row in annual.itertuples(index=False, name=None):
        worksheet.append(list(row))
    for row in worksheet.iter_rows(min_row=5, max_row=4 + len(annual)):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
        row[4].number_format = "0.00"
    add_positive_negative_rules(worksheet, f"E5:E{4 + len(annual)}")
    add_data_table(worksheet, 4, 4 + len(annual), len(headers), "AnnualStats")
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "年度涨跌幅（%）"
    chart.y_axis.title = "涨跌幅"
    chart.y_axis.majorGridlines = ChartLines()
    data = Reference(worksheet, min_col=5, min_row=4, max_row=4 + len(annual))
    categories = Reference(worksheet, min_col=1, min_row=5, max_row=4 + len(annual))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 7.5
    chart.width = 14
    chart.legend = None
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.series[0].graphicalProperties = GraphicalProperties(solidFill=RED)
    chart.series[0].graphicalProperties.line = LineProperties(solidFill=RED)
    worksheet.add_chart(chart, "A11")
    set_widths(worksheet, [12, 12, 12, 14, 14, 12, 12, 15, 14])
    worksheet.freeze_panes = "A5"


def build_monthly_sheet(workbook: Workbook, monthly: pd.DataFrame) -> None:
    worksheet = workbook.create_sheet("月度统计")
    worksheet.sheet_view.showGridLines = False
    worksheet.append(list(monthly.columns))
    style_header(worksheet[1])
    for row in monthly.itertuples(index=False, name=None):
        worksheet.append(list(row))
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")
        row[5].number_format = "0.00"
    add_positive_negative_rules(worksheet, f"F2:F{worksheet.max_row}")
    add_data_table(worksheet, 1, worksheet.max_row, len(monthly.columns), "MonthlyStats")
    add2 = ColorScaleRule(
        start_type="min",
        start_color="DDF2EA",
        mid_type="num",
        mid_value=0,
        mid_color="FFFFFF",
        end_type="max",
        end_color="FBE4E4",
    )
    worksheet.conditional_formatting.add(f"F2:F{worksheet.max_row}", add2)
    worksheet.freeze_panes = "B2"
    set_widths(worksheet, [12, 11, 11, 11, 12, 14, 13, 15, 14])


def build_metric_sheet(workbook: Workbook, summary: dict) -> None:
    worksheet = workbook.create_sheet("指标摘要")
    worksheet.sheet_view.showGridLines = False
    title_block(
        worksheet,
        "关键指标",
        "指标以 922 个交易日完整区间计算；用于描述历史风险与趋势。",
        width=4,
    )
    rows = [
        ("最新收盘", summary["keyMetrics"]["latestClose"], "元"),
        ("最新交易日", summary["keyMetrics"]["latestDate"], ""),
        ("区间收益", summary["keyMetrics"]["periodReturnPct"], "%"),
        ("年化波动率", summary["keyMetrics"]["annualizedVolatilityPct"], "%"),
        ("最大回撤", summary["keyMetrics"]["maxDrawdownPct"], "%"),
        ("回撤峰值日", summary["keyMetrics"]["maxDrawdownPeakDate"], ""),
        ("回撤低点日", summary["keyMetrics"]["maxDrawdownTroughDate"], ""),
        ("区间最高价", summary["keyMetrics"]["highestPrice"], "元"),
        ("最高价日期", summary["keyMetrics"]["highestDate"], ""),
        ("区间最低价", summary["keyMetrics"]["lowestPrice"], "元"),
        ("最低价日期", summary["keyMetrics"]["lowestDate"], ""),
        ("上涨交易日", summary["keyMetrics"]["positiveDays"], "日"),
        ("下跌交易日", summary["keyMetrics"]["negativeDays"], "日"),
        ("日胜率", summary["keyMetrics"]["winRatePct"], "%"),
        ("最新 MA20", summary["keyMetrics"]["latestMA20"], "元"),
        ("最新 MA60", summary["keyMetrics"]["latestMA60"], "元"),
        ("最新 MA120", summary["keyMetrics"]["latestMA120"], "元"),
        ("最新 RSI14", summary["keyMetrics"]["latestRSI14"], ""),
        ("最新换手率", summary["keyMetrics"]["latestTurnoverPct"], "%"),
    ]
    for row_index, header in enumerate(("指标", "数值", "单位"), start=4):
        worksheet.cell(row_index, 1, header)
    style_header(worksheet[4][:3])
    for row_index, values in enumerate(rows, start=5):
        for column_index, value in enumerate(values, start=1):
            cell = worksheet.cell(row_index, column_index, value)
            cell.alignment = Alignment(
                horizontal="left" if column_index == 1 else "right", vertical="center"
            )
            cell.fill = PatternFill(
                "solid", fgColor="FFFFFF" if row_index % 2 else "EEF0EB"
            )
            cell.border = Border(bottom=Side(style="hair", color=GRID))
    set_widths(worksheet, [24, 22, 12])
    worksheet.freeze_panes = "A5"


def build_chart_sheet(
    workbook: Workbook, frame: pd.DataFrame, sheet_name: str, kind: str
) -> None:
    worksheet = workbook.create_sheet(sheet_name)
    worksheet.sheet_view.showGridLines = False
    worksheet.append(
        ["日期", "收盘", "MA20", "MA60", "MA120"]
        if kind == "line"
        else ["日期", "成交量"]
    )
    style_header(worksheet[1])
    if kind == "line":
        for row in frame[["日期", "收盘", "MA20", "MA60", "MA120"]].itertuples(
            index=False, name=None
        ):
            worksheet.append(list(row))
        chart = LineChart()
        chart.title = "收盘价与移动平均线"
        chart.y_axis.title = "后复权价格（元）"
        chart.x_axis.title = "交易日"
        chart.style = 2
        chart.height = 10
        chart.width = 22
        chart.x_axis.tickLblSkip = 63
        chart.x_axis.tickMarkSkip = 63
        data = Reference(worksheet, min_col=2, max_col=5, min_row=1, max_row=worksheet.max_row)
        categories = Reference(
            worksheet, min_col=1, min_row=2, max_row=worksheet.max_row
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        for series, color in zip(chart.series, [INK, BLUE, GOLD, "7C5D9E"]):
            series.graphicalProperties.line = LineProperties(
                solidFill=color, w=20000
            )
        worksheet.add_chart(chart, "G2")
    else:
        for row in frame[["日期", "成交量"]].itertuples(index=False, name=None):
            worksheet.append(list(row))
        chart = BarChart()
        chart.title = "日成交量"
        chart.y_axis.title = "成交量（手）"
        chart.x_axis.title = "交易日"
        chart.style = 10
        chart.height = 9
        chart.width = 22
        chart.x_axis.tickLblSkip = 63
        chart.x_axis.tickMarkSkip = 63
        data = Reference(worksheet, min_col=2, min_row=1, max_row=worksheet.max_row)
        categories = Reference(
            worksheet, min_col=1, min_row=2, max_row=worksheet.max_row
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.legend = None
        chart.series[0].graphicalProperties = GraphicalProperties(solidFill=BLUE)
        chart.series[0].graphicalProperties.line = LineProperties(solidFill=BLUE)
        worksheet.add_chart(chart, "G2")
    set_widths(worksheet, [12, 12, 12, 12, 12])
    worksheet.freeze_panes = "B2"


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    market = pd.read_csv(DATA_DIR / "比亚迪_后复权日行情_2023-2026.csv")
    annual = pd.read_csv(DATA_DIR / "年度统计.csv")
    monthly = pd.read_csv(DATA_DIR / "月度统计.csv")
    summary = json.loads((DATA_DIR / "analysis-summary.json").read_text("utf-8"))
    market["日期"] = pd.to_datetime(market["日期"])

    workbook = Workbook()
    workbook.properties.title = "比亚迪股票价格走势分析与可视化（2023-2026）"
    workbook.properties.subject = "教材 P92 实战演练"
    workbook.properties.creator = "Codex Agent"
    workbook.properties.keywords = "比亚迪, 002594, 后复权, 股票分析"
    build_notes(workbook, summary)
    build_market_sheet(workbook, market)
    build_annual_sheet(workbook, annual)
    build_monthly_sheet(workbook, monthly)
    build_metric_sheet(workbook, summary)
    build_chart_sheet(workbook, market, "收盘走势", "line")
    build_chart_sheet(workbook, market, "成交量", "volume")

    for worksheet in workbook.worksheets:
        worksheet.sheet_properties.tabColor = INK
        worksheet.sheet_view.zoomScale = 90
    workbook.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
