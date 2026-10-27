from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "比亚迪_后复权日行情_2023-2026.csv"
SUMMARY = ROOT / "data" / "analysis-summary.json"
VALIDATION = ROOT / "data" / "validation.json"
WORKBOOK = ROOT / "deliverables" / "比亚迪股票价格走势分析_2023-2026.xlsx"


def check(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def main() -> None:
    frame = pd.read_csv(DATA)
    summary = json.loads(SUMMARY.read_text("utf-8"))
    validation = json.loads(VALIDATION.read_text("utf-8"))

    check(len(frame) == 922, f"Expected 922 rows, found {len(frame)}")
    check(frame["日期"].is_monotonic_increasing, "Dates are not sorted")
    check(not frame["日期"].duplicated().any(), "Duplicate trading dates found")
    check(frame.iloc[0]["日期"] == "2023-01-03", "Unexpected first date")
    check(frame.iloc[-1]["日期"] == "2026-10-27", "Unexpected last date")
    check(
        (frame["最高"] + 1e-6 >= frame[["开盘", "收盘", "最低"]].max(axis=1)).all(),
        "High price is below another OHLC value",
    )
    check(
        (frame["最低"] - 1e-6 <= frame[["开盘", "收盘", "最高"]].min(axis=1)).all(),
        "Low price is above another OHLC value",
    )
    check((frame[["开盘", "收盘", "最高", "最低", "成交量"]] > 0).all().all(), "Non-positive market value found")
    check(
        validation["overlapRows"] == 242
        and validation["adjustmentMatched"]
        and validation["sameTradingCalendar"],
        "2023 overlap validation failed",
    )
    check(
        all(value == 0 for value in validation["maxAbsDifferenceByField"].values()),
        "2023 field-level overlap difference is non-zero",
    )

    key = summary["keyMetrics"]
    check(key["latestClose"] == round(float(frame.iloc[-1]["收盘"]), 2), "Latest close mismatch")
    check(abs(key["maxDrawdownPct"] - float(frame["回撤"].min() * 100)) < 0.011, "Drawdown mismatch")
    check(
        abs(
            key["periodReturnPct"]
            - (float(frame.iloc[-1]["收盘"]) / float(frame.iloc[0]["收盘"]) - 1) * 100
        )
        < 0.011,
        "Period return mismatch",
    )
    check(np.isfinite(frame["MA120"].dropna()).all(), "Invalid MA120 values")

    book = openpyxl.load_workbook(WORKBOOK, data_only=False)
    expected_sheets = {
        "任务说明",
        "完整行情",
        "年度统计",
        "月度统计",
        "指标摘要",
        "收盘走势",
        "成交量",
    }
    check(expected_sheets.issubset(book.sheetnames), "Workbook is missing required sheets")
    chart_count = sum(len(sheet._charts) for sheet in book.worksheets)
    check(chart_count >= 2, f"Expected at least two native Excel charts, found {chart_count}")
    check(book["完整行情"].max_row == 923, "Workbook market sheet row count mismatch")

    report = {
        "ok": True,
        "rows": len(frame),
        "dateRange": [frame.iloc[0]["日期"], frame.iloc[-1]["日期"]],
        "overlapRows": validation["overlapRows"],
        "maxOverlapDifference": max(validation["maxAbsDifferenceByField"].values()),
        "latestClose": key["latestClose"],
        "periodReturnPct": key["periodReturnPct"],
        "annualizedVolatilityPct": key["annualizedVolatilityPct"],
        "maxDrawdownPct": key["maxDrawdownPct"],
        "excelSheets": book.sheetnames,
        "excelCharts": chart_count,
    }
    (ROOT / "artifacts").mkdir(exist_ok=True)
    (ROOT / "artifacts" / "data-validation.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
